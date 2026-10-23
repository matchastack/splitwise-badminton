import subprocess
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest
from telegram import Bot
from telegram.error import BadRequest

from badminton.bot.app import build_application
from badminton.bot.commands import ADMIN_COMMANDS, MEMBER_COMMANDS
from badminton.db.models import Member
from badminton.domain.sheet import SheetError
from badminton.services.spreadsheet import ImportResult, Match
from scripts.import_spreadsheet import announcement, parse_args, read_workbook, report
from scripts.setup_telegram import configure, webhook_url
from tests.conftest import TEST_SETTINGS
from tests.fake_telegram import ADMIN_ID, GROUP_ID, FakeTelegram

ROOT = Path(__file__).resolve().parents[1]


def command_names(commands: object) -> list[str]:
    return [command["command"] for command in commands]  # type: ignore[attr-defined]


async def test_configure_sets_webhook_and_command_menus() -> None:
    telegram = FakeTelegram()
    async with Bot(TEST_SETTINGS.bot_token, request=telegram) as bot:
        await configure(
            bot, TEST_SETTINGS, "https://example.vercel.app/", drop_pending_updates=True
        )

    webhook = telegram.last("setWebhook").params
    assert webhook["url"] == "https://example.vercel.app/api/telegram"
    assert webhook["secret_token"] == TEST_SETTINGS.webhook_secret
    assert webhook["allowed_updates"] == ["message", "callback_query"]
    assert webhook["drop_pending_updates"] is True

    menus = [
        (c.params["scope"], command_names(c.params["commands"]))
        for c in telegram.calls_to("setMyCommands")
    ]
    members = [c.command for c in MEMBER_COMMANDS]
    admins = [c.command for c in ADMIN_COMMANDS]
    assert menus == [
        ({"type": "default"}, members),
        ({"type": "chat", "chat_id": ADMIN_ID}, admins),
        ({"type": "chat_member", "chat_id": GROUP_ID, "user_id": ADMIN_ID}, admins),
    ]


async def test_configure_reports_an_admin_who_has_not_started_the_bot() -> None:
    telegram = FakeTelegram()
    telegram.fail("setMyCommands", 400, "Bad Request: chat not found", chat_id=ADMIN_ID)
    async with Bot(TEST_SETTINGS.bot_token, request=telegram) as bot:
        refused = await configure(bot, TEST_SETTINGS, "https://example.vercel.app")

    assert refused == [f"admin {ADMIN_ID}, private chat: Chat not found"]
    assert telegram.last("setWebhook").params["url"] == "https://example.vercel.app/api/telegram"
    scopes = [c.params["scope"] for c in telegram.calls_to("setMyCommands")]
    assert {"type": "chat_member", "chat_id": GROUP_ID, "user_id": ADMIN_ID} in scopes


async def test_configure_still_raises_when_the_webhook_is_refused() -> None:
    telegram = FakeTelegram()
    telegram.fail("setWebhook", 400, "Bad Request: bad webhook: HTTPS url must be provided")
    async with Bot(TEST_SETTINGS.bot_token, request=telegram) as bot:
        with pytest.raises(BadRequest, match=r"(?i)bad webhook"):
            await configure(bot, TEST_SETTINGS, "https://example.vercel.app")


def test_webhook_url_requires_https() -> None:
    with pytest.raises(ValueError, match="https"):
        webhook_url("http://example.com")


def test_polling_application_keeps_its_updater() -> None:
    application = build_application(TEST_SETTINGS, None, polling=True)  # type: ignore[arg-type]

    assert application.updater is not None
    assert build_application(TEST_SETTINGS, None).updater is None  # type: ignore[arg-type]


def write_sheet(path: Path) -> Path:
    workbook = openpyxl.Workbook()
    items = workbook.active
    items.title = "Items"
    rows = [
        ("Date", "Item", "Price", "Paid By", "zy", "yt", "Number of Shares", "Price per Share"),
        (datetime(2026, 10, 18), "court", 14.8, "zy", 1, 1, "=SUM(E2:F2)", "=C2/G2"),
        (datetime(2026, 10, 18), "shuttle", 76.79, "yt", 0, None, "=SUM(E3:F3)", "=C3/G3"),
    ]
    for row in rows:
        items.append(row)
    workbook.create_sheet("Total")
    workbook.save(path)
    return path


def test_import_reads_the_items_sheet_and_reports_the_tally(tmp_path: Path) -> None:
    summary = read_workbook(write_sheet(tmp_path / "Badminton.xlsx"))
    zy = Member(id=1, display_name="Zhen Yi", username="zzyi", telegram_user_id=111)

    lines = report(summary, [Match("zy", zy, "same name"), Match("yt", None, "new guest")], [zy])

    assert lines[:6] == [
        "Bot members now:",
        "  #1 Zhen Yi (@zzyi)",
        "",
        "Player columns:",
        "  zy       → Zhen Yi (@zzyi), same name",
        "  yt       → new guest member",
    ]
    assert lines[7:11] == [
        "Player         Paid        Spend    Surplus     Tube   Opening",
        "zy            14.80       7.4000     7.4000              +7.40",
        "yt            76.79       7.4000    69.3900   -76.79     -7.40",
        "Total         91.59      14.8000    76.7900   -76.79     +0.00",
    ]
    assert lines[-1].startswith("Row 3 is a tube nobody has used yet: $76.79 bought by yt.")


def test_import_needs_an_items_sheet(tmp_path: Path) -> None:
    path = tmp_path / "Other.xlsx"
    openpyxl.Workbook().save(path)

    with pytest.raises(SheetError, match="no 'Items' sheet"):
        read_workbook(path)


def test_import_announces_each_carried_over_balance(tmp_path: Path) -> None:
    summary = read_workbook(write_sheet(tmp_path / "Badminton.xlsx"))
    result = ImportResult(
        member_ids={"zy": 1, "yt": 2},
        guests=("yt",),
        opening_cents={"zy": 740, "yt": -740},
        batch_id=1,
        balances={1: 740, 2: -740},
    )

    text = announcement(summary, result, {1: "Zhen <Yi>", 2: "yt"})

    assert text == (
        "📒 <b>Balances carried over from the spreadsheet</b>, up to Sun 18 Oct:\n"
        "• Zhen &lt;Yi&gt; +$7.40\n"
        "• yt −$7.40\n"
        "+ means the group owes them, − means they owe the group.\n"
        "\n"
        "🏸 The $76.79 tube yt bought is shuttle batch #1. "
        "It's billed by hours played once the next one is bought."
    )


def test_import_member_choices_are_name_equals_id() -> None:
    assert parse_args(["x.xlsx", "--member", "zy=3"]).member == [("zy", 3)]
    with pytest.raises(SystemExit):
        parse_args(["x.xlsx", "--member", "zy"])


@pytest.mark.parametrize(
    ("script", "flag"),
    [
        ("setup_telegram.py", "--base-url"),
        ("dev_polling.py", "--tick-interval"),
        ("import_spreadsheet.py", "--apply"),
    ],
)
def test_scripts_document_their_arguments(script: str, flag: str) -> None:
    result = subprocess.run(
        [sys.executable, str(ROOT / "scripts" / script), "--help"],
        capture_output=True,
        text=True,
        check=True,
    )

    assert flag in result.stdout
