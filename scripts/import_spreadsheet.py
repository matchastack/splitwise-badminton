"""Carry the balances in Badminton.xlsx over to the bot (FR-MIG-1 to FR-MIG-4).

By default this is a dry run: it shows which member each player column becomes and
how every balance tallies, and saves nothing. Run the database migrations first:

    DATABASE_URL='postgresql://…' uv run alembic upgrade head
    uv run --env-file .env scripts/import_spreadsheet.py Badminton.xlsx
    uv run --env-file .env scripts/import_spreadsheet.py Badminton.xlsx --member alice=1 --apply

A player column matches the member with the same name or @username; anyone else is
added as a guest, who the admin links to their Telegram account from /members later.
"""

from __future__ import annotations

import argparse
import asyncio
import html
import sys
from collections.abc import Sequence
from fractions import Fraction
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl
from telegram import Bot
from telegram.constants import ParseMode
from telegram.error import TelegramError

from badminton.config import Settings
from badminton.db.engine import create_engine, create_sessionmaker
from badminton.db.models import Member
from badminton.domain.money import fmt_sgd, fmt_signed
from badminton.domain.sheet import SheetError, SheetSummary, read_items
from badminton.domain.timeutil import fmt_date, sgt_date, utc_now
from badminton.services.members import all_members, member_label
from badminton.services.notify import announce
from badminton.services.spreadsheet import (
    ImportRefused,
    ImportResult,
    Match,
    TallyMismatch,
    check_can_import,
    cutoff,
    import_sheet,
    match_members,
)

ITEMS_SHEET = "Items"


def read_workbook(path: Path) -> SheetSummary:
    """The ``Items`` sheet's summary. Formula cells are read as Excel last saved them."""
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        if ITEMS_SHEET not in workbook.sheetnames:
            raise SheetError(f"The workbook has no {ITEMS_SHEET!r} sheet.")
        rows = workbook[ITEMS_SHEET].iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            raise SheetError(f"The {ITEMS_SHEET!r} sheet is empty.")
        return read_items(header, enumerate(rows, start=2))
    finally:
        workbook.close()


def report(summary: SheetSummary, matches: Sequence[Match], members: Sequence[Member]) -> list[str]:
    """Who each player becomes, and how their balance tallies with the sheet."""
    lines = ["Bot members now:"]
    lines += [f"  #{m.id} {member_label(m)}" for m in members] or ["  (none yet)"]
    lines += ["", "Player columns:"]
    for match in matches:
        target = (
            "new guest member"
            if match.member is None
            else f"{member_label(match.member)}, {match.how}"
        )
        lines.append(f"  {match.name:<8} → {target}")

    opening = summary.opening_cents()
    lines += ["", _row("Player", "Paid", "Spend", "Surplus", "Tube", "Opening")]
    for name in summary.players:
        tube = summary.pending_cents(name)
        lines.append(
            _row(
                name,
                _amount(summary.paid[name], 2),
                _amount(summary.spend[name], 4),
                _amount(summary.surplus(name), 4),
                _amount(Fraction(-tube, 100), 2) if tube else "",
                _amount(Fraction(opening[name], 100), 2, signed=True),
            )
        )
    tube_total = summary.pending.amount_cents if summary.pending else 0
    lines.append(
        _row(
            "Total",
            _amount(sum(summary.paid.values(), Fraction(0)), 2),
            _amount(sum(summary.spend.values(), Fraction(0)), 4),
            _amount(sum(map(summary.surplus, summary.players), Fraction(0)), 4),
            _amount(Fraction(-tube_total, 100), 2),
            _amount(Fraction(sum(opening.values()), 100), 2, signed=True),
        )
    )
    lines += [
        "",
        "Surplus = Paid − Spend, as on the Total sheet. Opening = Surplus − Tube, in cents.",
    ]
    if summary.pending is not None:
        pending = summary.pending
        lines.append(
            f"Row {pending.row} is a tube nobody has used yet: {fmt_sgd(pending.amount_cents)} "
            f"bought by {pending.payer}. It becomes the open shuttle batch from "
            f"{fmt_date(sgt_date(cutoff(summary)))} and is billed by hours played once the "
            "next tube is recorded with /shuttle."
        )
    return lines


def announcement(summary: SheetSummary, result: ImportResult, names: dict[int, str]) -> str:
    """The group post (NFR-6): every carried-over balance, and the open batch."""
    lines = [
        f"📒 <b>Balances carried over from the spreadsheet</b>, up to {fmt_date(summary.last_day)}:"
    ]
    for player in summary.players:
        name = html.escape(names[result.member_ids[player]])
        lines.append(f"• {name} {fmt_signed(result.opening_cents[player])}")
    lines.append("+ means the group owes them, − means they owe the group.")
    if summary.pending is not None and result.batch_id is not None:
        payer = html.escape(names[result.member_ids[summary.pending.payer]])
        lines += [
            "",
            f"🏸 The {fmt_sgd(summary.pending.amount_cents)} tube {payer} bought is shuttle "
            f"batch #{result.batch_id}. It's billed by hours played once the next one is bought.",
        ]
    return "\n".join(lines)


def _row(*cells: str) -> str:
    name, paid, spend, surplus, tube, opening = cells
    return f"{name:<8} {paid:>10} {spend:>12} {surplus:>10} {tube:>8} {opening:>9}"


def _amount(value: Fraction, places: int, *, signed: bool = False) -> str:
    return f"{float(value):{'+' if signed else ''},.{places}f}"


def member_choice(text: str) -> tuple[str, int]:
    name, sep, member_id = text.partition("=")
    if not sep or not name.strip() or not member_id.strip().isdigit():
        raise argparse.ArgumentTypeError(f"expected NAME=MEMBER_ID, e.g. alice=1, not {text!r}")
    return name.strip(), int(member_id)


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawTextHelpFormatter
    )
    parser.add_argument("path", type=Path, help="The spreadsheet, e.g. Badminton.xlsx")
    parser.add_argument(
        "--member",
        type=member_choice,
        action="append",
        default=[],
        metavar="NAME=ID",
        help="Use member ID for the player column NAME (IDs are listed in the dry run)",
    )
    parser.add_argument("--apply", action="store_true", help="Save the import (default: dry run)")
    parser.add_argument(
        "--no-announce", action="store_true", help="Don't post the balances in the group"
    )
    return parser.parse_args(argv)


async def main(argv: list[str] | None = None) -> None:
    args = parse_args(argv)
    summary = read_workbook(args.path)
    settings = Settings.from_env()
    engine = create_engine(settings.database_url)
    try:
        async with create_sessionmaker(engine)() as db:
            await check_can_import(db, summary)  # so a dry run is refused like --apply
            matches = await match_members(db, summary.players, dict(args.member))
            print("\n".join(report(summary, matches, await all_members(db))))
            if not args.apply:
                print("\nDry run: nothing was saved. Run again with --apply to import.")
                return
            result = await import_sheet(db, summary, matches, now=utc_now())
            await db.commit()
            names = {member.id: member.display_name for member in await all_members(db)}
    finally:
        await engine.dispose()

    print("\n✅ Imported. Each balance moved by its Opening amount, and all balances sum to $0.00.")
    if result.guests:
        print(
            f"Added as guests: {', '.join(result.guests)}. Link them from /members when they join."
        )
    if args.no_announce:
        return
    if settings.group_chat_id is None:
        print("GROUP_CHAT_ID isn't set, so the balances weren't posted in the group.")
        return
    try:
        async with Bot(settings.bot_token) as bot:
            # A plain Bot has no default parse mode; the bot app sets HTML for all its messages.
            text = announcement(summary, result, names)
            await announce(bot, settings, text, parse_mode=ParseMode.HTML)
    except TelegramError as error:
        print(f"The import is saved, but posting in the group failed: {error}", file=sys.stderr)
        return
    print("Posted the balances in the group.")


if __name__ == "__main__":
    try:
        asyncio.run(main())
    except (SheetError, ImportRefused, TallyMismatch) as error:
        sys.exit(f"Not imported: {error}")
